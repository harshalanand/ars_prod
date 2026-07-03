"""Generate RULE_MASTER.docx and RULE_MASTER.xlsx from the rule data below.

Run:  backend\venv\Scripts\python.exe docs\build_rule_master.py
Outputs land next to this file as RULE_MASTER.docx and RULE_MASTER.xlsx.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.shared import Pt, RGBColor

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


# ─── Rule data ────────────────────────────────────────────────────────────────
# Each section: title, body paragraphs, optional list of (table_title, headers, rows).

INVARIANTS = [
    ("1", "OPT uniqueness", "One (WERKS, MAJ_CAT, GEN_ART, CLR) ⇒ exactly one OPT_TYPE (RL ∨ TBC ∨ TBL ∨ MIX) after Listing Part 3.6", "Hold/unpark paths if OPT_TYPE changes silently"),
    ("2", "Growth at MJ + grid only", "Growth % is applied at MAJ_CAT × grid level. Never per OPT_TYPE", "Any code that scales an OPT-row's MBQ by growth"),
    ("3", "MBQ sparseness", "MBQ = 0 means no constraint at that grain. Do not apply 1.30× breach when MBQ = 0", "Sec-cap path if it treats 0 as 'zero allowed'"),
    ("4", "Sec-cap dims must propagate", "FAB, MACRO_MVGR, MICRO_MVGR, M_VND_CD, RNG_SEG must survive listing → listed → alloc", "Any pipeline stage that drops a column"),
    ("5", "ACS_D ≠ daily sale", "ACS_D is an aggregated demand figure. For velocity, use MAX_DAILY_SALE", "Anywhere ACS_D is divided by days as if it were a rate"),
    ("6", "RNG_SEG = MRP tier", "Values are E / V / P / SP (essential / value / premium / super-premium). Not a free text field", "UI filters that allow arbitrary values"),
]

MSA_STEPS = [
    ("1",  "Filter SLOC",           "Keep only rows where SLOC IN (selected_slocs). Out-of-scope shelves never contribute stock."),
    ("2",  "Normalize numerics",    "Cast to numeric, NaN → 0"),
    ("3",  "Fill missing dims",     "Apply column defaults (e.g. CLR='NA')"),
    ("4",  "SEG filter",            "Keep SEG IN ('APP','GM')"),
    ("5",  "Pivot by SLOC",         "Rename ST_CD → RDC immediately after pivot (must happen before any downstream column lookup)"),
    ("6",  "Universe backfill",     "_load_universe(slocs, date) returns (RDC, GEN_ART) union of: stock in selected SLOCs, open ARS_PEND_ALC, open ARS_NL_TBL_HOLD_TRACKING. Backfill VAR_ARTs from vw_master_product so every PEND/HOLD has a row."),
    ("7",  "Merge PEND",            "ARS_PEND_ALC → PEND_QTY on (RDC, ARTICLE_NUMBER)"),
    ("8",  "Merge HOLD",            "ARS_NL_TBL_HOLD_TRACKING → HOLD_QTY (map WERKS → RDC)"),
    ("9",  "Compute FNL_Q",         "FNL_Q = max(STK − PEND − HOLD, 0)"),
    ("10", "Threshold",             "Keep groups where Σ FNL_Q + Σ PEND_QTY + Σ HOLD_QTY > threshold (admits pend-only / hold-only groups)"),
    ("11", "Aggregate to GEN_ART",  "Group VAR_ART → GEN_ART by (RDC, MAJ_CAT, GEN_ART_NUMBER, CLR)"),
]

MSA_RECON = [
    ("R1", "Σ TOTAL.PEND_QTY == Σ ARS_PEND_ALC.PEND_QTY for IS_CLOSED=0",                                              "Δ = 0"),
    ("R2", "Σ TOTAL.HOLD_QTY == Σ ARS_NL_TBL_HOLD_TRACKING.HOLD_REM for IS_CLOSED=0 (via WERKS→RDC)",                  "Δ = 0"),
    ("R3", "Σ TOTAL.STK_QTY == Σ VW_ET_MSA_STK_WITH_MASTER.STK_Q for run date, selected SLOCs, SEG IN ('APP','GM')",  "Δ = 0"),
    ("R4", "count(distinct VAR_ART.ARTICLE_NUMBER) == count(distinct TOTAL.ARTICLE_NUMBER) per passing group",        "Equal"),
    ("R5", "Per (RDC, MAJ_CAT, GEN_ART, CLR): GEN_ART.X == Σ VAR_ART.X for X ∈ {STK_QTY, PEND_QTY, HOLD_QTY, FNL_Q}",  "Equal"),
]

MSA_WRITE_PATHS = [
    ("1", "msa_result_storage.store_results",        "MSA Generate",                                                       "TRUNCATE+INSERT all 3 tables; auto-runs paths 7+8"),
    ("2", "adjust_msa_after_pend_insert",            "PEND_ALC INSERT (manual, CSV, approve_parked)",                      "PEND_QTY/FNL_Q for affected (RDC, ART)"),
    ("3", "apply_pend_alc_delta(sign=±1)",           "PEND_ALC mutation with explicit sign",                                "PEND_QTY/FNL_Q delta in all 3"),
    ("4", "apply_pend_alc_delta_by_session",         "Approve Parked / session-scoped revert",                             "Same as #3, scoped by SESSION_ID"),
    ("5", "apply_hold_clear",                        "Hold Dashboard 'Clear'",                                              "HOLD_QTY/FNL_Q for closed hold rows"),
    ("6", "apply_hold_revise",                       "Hold Dashboard 'Revise'",                                             "HOLD_QTY/FNL_Q for revised hold rows"),
    ("7", "bootstrap_msa_pend_sync",                 "End of #1, post-revert, post-grid-build",                             "Full reseed of PEND_QTY/FNL_Q"),
    ("8", "bootstrap_msa_hold_sync",                 "End of #1, after approve_parked PEND delta",                          "Reseed HOLD_QTY/FNL_Q"),
]

GRID_TAXONOMY = [
    ("MJ_RNG_SEG",      "Primary", "MAJ_CAT × RNG_SEG",     "Main MBQ & growth driver"),
    ("MJ_FAB",          "Sec-cap", "MAJ_CAT × FAB",          "Fabric breach gate"),
    ("MJ_MICRO_MVGR",   "Sec-cap", "MAJ_CAT × MICRO_MVGR",   "Micro-vendor-group gate"),
    ("MJ_<other>",      "Sec-cap", "MAJ_CAT × <dim>",        "Configurable extras"),
]

GRID_SCHEMA_RULES = [
    ("_ensure_hierarchy_table is add-only",                                            "Deleting a grid / setting Inactive does NOT drop a column from ARS_GRID_HIERARCHY; data survives accidental deletes"),
    ("Physical column order is no longer reshuffled",                                  "Consumers reference by name, not ordinal"),
    ("Orphan column removal goes through POST /grid-builder/hierarchy/compact",        "Explicit admin action; default dry_run=true. Includes MERGE_<X> orphans when parent X is gone"),
    ("No DBCC SHRINKFILE after Run-All",                                               "Routine shrink fragments indexes; file regrows next run anyway"),
    ("S5 _insert_missing_msa_rows skipped when grid.pivot_only=1",                     "Article-grain grids would only get (NA,NA,NA) placeholders polluting downstream"),
    ("bootstrap_msa_pend_sync at end of Run-All is kept",                              "Cheap safety net; prevents PEND/FNL_Q drift after heavy rebuild"),
]

OPT_TYPE_TABLE = [
    ("1", "MIX (a) — nothing to send",                       "NOT ADEQUATE_STK AND MSA_FNL_Q = 0 AND RL_HOLD_QTY = 0",                                                    "MIX"),
    ("2", "MIX (b) — poor color fill",                        "VAR_COUNT > 0 AND (VAR_FNL_COUNT / VAR_COUNT < threshold OR VAR_FNL_COUNT < min_size_count)",              "MIX"),
    ("3", "RL — adequate stock or live hold, plus fresh MSA", "(ADEQUATE_STK OR RL_HOLD_QTY > 0) AND MSA_FNL_Q > 0",                                                       "RL"),
    ("4", "TBC — low (positive) stock, supply available",     "0 < STK_TTL < threshold × ACS_D AND (MSA_FNL_Q > 0 OR RL_HOLD_QTY > 0)",                                    "TBC"),
    ("5", "TBL — zero/negative stock, supply available",      "STK_TTL ≤ 0 AND (MSA_FNL_Q > 0 OR RL_HOLD_QTY > 0)",                                                        "TBL"),
    ("6", "Catch-all",                                        "(none of the above)",                                                                                       "MIX"),
]

MIX_MODES = [
    ("st_maj_rng (default)",   "1 MIX row per (WERKS, MAJ_CAT, RNG_SEG) — finer"),
    ("st_maj",                  "1 MIX row per (WERKS, MAJ_CAT)"),
    ("each",                    "Keep each MIX row as-is (tag only)"),
]

OPT_STATUS_TABLE = [
    ("RL",   "any",              "RL"),
    ("TBC",  "> 0 (regular)",    "TBC"),
    ("TBC",  "> 0 (special)",    "(other branch — see code)"),
    ("TBC",  "= 0",              "MIX"),
    ("TBL",  "> 0",              "TBL (or branch)"),
    ("TBL",  "= 0",              "TBL"),
    ("else", "—",                "ISNULL(OPT_TYPE, 'MIX')"),
]

ALLOC_MODES = [
    ("pandas (default)",  "Vectorized cumulative-window race"),
    ("per_opt",            "One-OPT-at-a-time sequential engine; flips ARS_PER_OPT_MODE=1"),
    ("sequential",         "Single-thread SQL fallback"),
]

FNL_Q_REM_TABLE = [
    ("0",   "PAK_SZ_ROUND(..., short=stock=N)", "Pool exhausted"),
    ("> 0", "PAK_SZ_GATE(req=R, pak=P)",        "Pak rule fired (stock available but not pak-aligned)"),
    ("> 0", "from_hold=N in round 2+",           "Reservation replay (correct on RL/TBC)"),
]

HOLD_PEND_RISKS = [
    ("ARS_PEND_ALC has only PK on ID — no unique on (SESSION_ID, RDC, ST_CD, ARTICLE_NUMBER, ALLOC_MODE)", "pend_alc_service.py:16",        "Duplicate manual rows silently double-count after apply_pend_alc_delta aggregates"),
    ("Manual upload has no required Pydantic validation",                                                  "pend_alc.py:2222",              "Blank rdc / orphan articles get inserted → MSA delta matches 0 rows → MSA stays stale until next bootstrap"),
    ("apply_adhoc_close does not sync MSA",                                                                "pend_alc_service.py:2878",      "MSA_TOTAL.PEND_QTY stays inflated until next full Generate; listings see lower FNL_Q than reality"),
    ("Adhoc close with blank ST_CD is 'any-store' wildcard",                                               "pend_alc_service.py:2961",      "Empty Excel ST_CD column closes every open row for (RDC, ARTICLE) across all stores"),
    ("apply_do_deductions FIFO partition does not include ALLOC_MODE",                                     "pend_alc_service.py:3371,3376", "RL DO can settle a TBC allocation (and vice versa) — cross-attribution risk"),
    ("update_bdc_history_with_do do_qty=0 cancellation has wildcard scope",                                 "pend_alc_service.py:2772",      "Cancel with only (rdc, art) + do_qty=0 flips every OPEN BDC for that pair"),
    ("_check_adhoc_close_revert gates on LAST_DO_AT/LAST_BDC_AT but apply_adhoc_close doesn't touch them", "pend_alc_service.py:3011",      "Revert always passes; if BDC was re-generated, two OPEN history rows can co-exist for same (RDC, ST_CD, ART)"),
]

CHECKLIST = [
    # (section, item)
    ("MSA",       "R1: Σ TOTAL.PEND == Σ PEND_ALC (open)"),
    ("MSA",       "R2: Σ TOTAL.HOLD == Σ HOLD_TRACKING (open) with WERKS→RDC"),
    ("MSA",       "R3: Σ TOTAL.STK == Σ stock view for run date + selected SLOCs + APP/GM"),
    ("MSA",       "R4: Per group, VAR_ART article-count = TOTAL article-count"),
    ("MSA",       "R5: Per (RDC, MAJ_CAT, GEN_ART, CLR), GEN_ART.X = Σ VAR_ART.X for X in {STK, PEND, HOLD, FNL_Q}"),
    ("Grid",      "Every grid row carries FAB, MACRO_MVGR, MICRO_MVGR, M_VND_CD, RNG_SEG"),
    ("Grid",      "Growth % applied at MAJ_CAT × grid, not at OPT_TYPE level"),
    ("Grid",      "For each grid row with MBQ=0, the 1.30× breach was skipped"),
    ("Grid",      "Run-All was followed by bootstrap_msa_pend_sync"),
    ("Listing",   "Each row has exactly one OPT_TYPE after Part 3.6"),
    ("Listing",   "For each MIX row, matches one branch in the OPT_TYPE decision table"),
    ("Listing",   "MIX aggregation produced expected granularity per mix_mode"),
    ("Listing",   "Per (WERKS, MAJ_CAT), Σ SHIP_QTY per OPT_TYPE ≤ cap_pct × MJ_REQ"),
    ("Listing",   "Per OPT_TYPE, Σ SHIP_QTY ≤ mbq_cap_pct × MJ_MBQ_ORIG"),
    ("Listing",   "HOLD_DAYS extension applied only to TBL rows"),
    ("Allocation","Per OPT, per size: SHIP_QTY ≤ I_ROD × SZ_MBQ − SZ_STK + (pak−1)"),
    ("Allocation","RL/TBC rows with RL_HOLD_QTY > 0: hold consumed first, then pool"),
    ("Allocation","FNL_Q_REM consistent with ALLOC_REMARKS"),
    ("Allocation","OPT_STATUS matches Part 8.5 table"),
    ("Allocation","TBL_LISTED_DATE populated for shipped TBL rows"),
    ("Hold/Pend", "Every parked_history transition has a matching row"),
    ("Hold/Pend", "No reverted allocation left stale PEND_QTY in MSA"),
    ("Hold/Pend", "No adhoc-close run since last MSA Generate (else suspect inflated PEND)"),
]

GAPS = [
    ("9.1", "Listing Part 3.6 — MIX(b) 'poor color fill' branch contradicts simplification target",
            "listing.py:1387-1390 still tags MIX based on VAR_FNL_COUNT/VAR_COUNT < threshold even when MSA_FNL_Q > 0. Intent (clarified 2026-06-23): MIX is supply-only (MSA_FNL_Q=0 AND RL_HOLD_QTY=0). Symptom: a run produced 397 aggregate MIX lines, 396 with live supply — those OPTs were short-circuited and never shipped.",
            "Split MIX(b) into a separate POOR_COLOR_FILL flag column. Keep the OPT in RL/TBC/TBL based on stock; surface fill quality for review."),
    ("9.2", "Per-OPT exec_order=round_first is configured but not wired",
            "listing.py:142-144 accepts round_first (R1 across all OPT_TYPEs → R2 → …), advertised as 'fairer to TBL', but the engine maps both options to opt_type_first.",
            "Either wire round_first end-to-end or drop the option from UI / config. Today it silently no-ops."),
    ("9.3", "apply_adhoc_close doesn't sync MSA — silent data drift",
            "Closes PEND rows but doesn't invoke apply_pend_alc_delta(sign=-1) or bootstrap_msa_pend_sync. Listings between adhoc close and next MSA Generate see lower FNL_Q than reality.",
            "End apply_adhoc_close with a session-scoped apply_pend_alc_delta(sign=-1). Or gate listings from running between adhoc close and next Generate."),
    ("9.4", "Adhoc close + DO cancel — wildcard ST_CD semantics are unsafe defaults",
            "Blank ST_CD = 'any store'; do_qty=0 with only (rdc, art) = 'every open BDC'. Copy-paste from Excel commonly leaves these blank.",
            "Require explicit ALL_STORES=true in the payload to invoke the wildcard. Default blank ST_CD to a rejection with a confirmation prompt."),
    ("9.5", "ARS_PEND_ALC grain claim vs reality",
            "Docstring claims (SESSION_ID, RDC, ST_CD, ARTICLE_NUMBER, ALLOC_MODE) is the grain, but only PK exists; write_manual_pend_alc doesn't enforce.",
            "Either add a unique index matching the documented grain, or update the docstring. Current mismatch enables duplicate-row double-counting."),
    ("9.6", "DO FIFO partition ignores ALLOC_MODE",
            "apply_do_deductions partitions FIFO by (RDC, ST_CD, ART) without ALLOC_MODE. An RL DO can consume a TBC PEND row and vice versa.",
            "Include ALLOC_MODE in the FIFO partition unless cross-attribution is intentional. If intentional, document it and surface in reporting."),
    ("9.7", "msa_result_storage.py docstring is wrong about DB target",
            "Says 'Main DB session, not Data DB' but all four MSA output tables live in the Data DB.",
            "Fix the docstring. Cheap, but prevents the next engineer from chasing a non-bug."),
    ("9.8", "MATNR / QTY legacy columns on ARS_PEND_ALC",
            "ARS_PEND_ALC carries MATNR (bigint) and QTY (int) that are never written or read; survived a schema migration.",
            "Drop them in a maintenance window (ALTER TABLE per column with backup). Until then, document them on the schema page as 'unused — do not select'."),
]


# ─── Word (.docx) generation ──────────────────────────────────────────────────

def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    return h


def add_para(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    return p


def add_bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light Grid Accent 1"
    # Header row
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for paragraph in hdr[i].paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(10)
    # Data rows
    for ri, row in enumerate(rows, start=1):
        cells = table.rows[ri].cells
        for ci, val in enumerate(row):
            cells[ci].text = str(val)
            for paragraph in cells[ci].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
            cells[ci].vertical_alignment = WD_ALIGN_VERTICAL.TOP
    return table


def build_docx(out_path: Path):
    doc = Document()

    # Title page-ish header
    add_heading(doc, "ARS Rule Master", level=0)
    add_para(doc, "MSA · Grid Builder · Listing · Allocation · Hold/Pending Allocation", italic=True)
    add_para(doc,
             "Audience: anyone reviewing an ARS run output. Use this as the single reference to answer "
             "'is the system following the rules?' and to suggest new rules where you see gaps.")
    add_para(doc,
             "Source files (deep dives): .claude/agents/ars_flow_kb/INDEX.md (per-module rule files); "
             "backend/app/services/ (implementation); backend/app/api/v1/endpoints/listing.py (OPT_TYPE classification).")
    add_para(doc, "Generated 2026-06-26 from docs/RULE_MASTER.md", italic=True)

    # 1. Pipeline
    add_heading(doc, "1. Pipeline at a glance", level=1)
    add_para(doc, "MSA Calculation → Grid Builder → Listing → Allocation → Hold/Pending Allocation")
    add_para(doc, "Outputs flow: ARS_MSA_* → ARS_GRID_* → ARS_LISTING → ARS_LISTED / ARS_ALLOC_HISTORY → ARS_PEND_ALC + ARS_NL_TBL_HOLD_TRACKING")

    # 2. Invariants
    add_heading(doc, "2. Cross-cutting invariants", level=1)
    add_para(doc, "These apply across MSA / Grid / Listing / Allocation. A run that violates any is broken.")
    add_table(doc, ["#", "Invariant", "Meaning", "Common violation site"], INVARIANTS)

    # 3. MSA
    add_heading(doc, "3. MSA Calculation rules", level=1)
    add_para(doc, "Source: backend/app/services/msa_service.py — MSAService.calculate()", italic=True)
    add_para(doc, "Trigger: POST /api/v1/msa/calculate (via msa_job_service)", italic=True)
    add_para(doc, "Outputs (rep_data DB): ARS_MSA_TOTAL, ARS_MSA_VAR_ART, ARS_MSA_GEN_ART, MSA_Calculation_Sequence", italic=True)

    add_heading(doc, "3.1 11-step MSA flow (universe-anchored, June 2026)", level=2)
    add_table(doc, ["Step", "Action", "Rule"], MSA_STEPS)

    add_heading(doc, "3.2 Universe rule (post-correction, June 2026)", level=2)
    add_para(doc, "Stock contribution stays SLOC-scoped. Products with stock only on non-selected shelves do not enter MSA. The universe expands only when a real obligation (PEND or HOLD) attaches to the (RDC, GEN_ART) pair. Cross-shelf stock alone is never enough.")

    add_heading(doc, "3.3 Reconciliation guarantees", level=2)
    add_para(doc, "All must hold after every Generate. If any ≠ 0, the bug is upstream of allocation.")
    add_table(doc, ["ID", "Check", "Pass condition"], MSA_RECON)

    add_heading(doc, "3.4 Eight write paths to ARS_MSA_TOTAL / VAR_ART / GEN_ART", level=2)
    add_para(doc, "Only path 1 inserts rows. Paths 2-8 are UPDATE-only and assume the row exists — that's why universe-anchored Step 6 matters.")
    add_table(doc, ["#", "Function", "Trigger", "Writes"], MSA_WRITE_PATHS)

    # 4. Grid Builder
    add_heading(doc, "4. Grid Builder rules", level=1)
    add_para(doc, "Source: backend/app/services/grid_calculations.py, backend/app/api/v1/endpoints/grid_builder.py", italic=True)

    add_heading(doc, "4.1 Grid taxonomy", level=2)
    add_table(doc, ["Grid", "Type", "Grain", "Purpose"], GRID_TAXONOMY)
    add_para(doc, "RNG_SEG values: E (essential) / V (value) / P (premium) / SP (super-premium).")

    add_heading(doc, "4.2 Required dimensions on every grid row", level=2)
    add_para(doc, "Every grid row must carry FAB, MACRO_MVGR, MICRO_MVGR, M_VND_CD, RNG_SEG. If any are dropped between listing → listed → alloc, sec-cap silently loses grids (invariant 4).")

    add_heading(doc, "4.3 Growth % rule", level=2)
    add_para(doc, "Growth is applied only at MAJ_CAT + grid level. Never per OPT_TYPE. Holds for the main path and any fallback path (invariant 2).")

    add_heading(doc, "4.4 Sec-cap dispatch (strict mode, default since 2026-06-16)", level=2)
    add_para(doc, "When apply_sec_cap_in_normal = True (default), for each grid where ARS_GRID_BUILDER.sec_cap_applicable = 1:")
    add_para(doc, "    budget = max(0, MBQ_ORIG × sec_cap_pct% − STK_TTL)")
    add_para(doc, "    breach ⇒ block the OPT row at the main-pass pre-gate")
    add_para(doc, "sec_cap_pct is per-grid (ARS_GRID_BUILDER.sec_cap_pct, default 130 when NULL). Strict binary semantics — no high-demand override.")

    add_heading(doc, "4.5 Grid schema management rules", level=2)
    add_table(doc, ["Rule", "Reason"], GRID_SCHEMA_RULES)

    # 5. Listing
    add_heading(doc, "5. Listing rules", level=1)
    add_para(doc, "Source: backend/app/api/v1/endpoints/listing.py (Parts 1 → 8.5)", italic=True)

    add_heading(doc, "5.1 OPT_TYPE classification (Part 3.6) — decision table", level=2)
    add_para(doc, "Evaluated top to bottom, first match wins. Parameters: stock_threshold_pct (default 0.6), default_acs_d (18), min_size_count (3).")
    add_para(doc, "Let SUPPLY_OK = (MSA_FNL_Q > 0) OR (RL_HOLD_QTY > 0)", italic=True)
    add_para(doc, "Let ADEQUATE_STK = STK_TTL ≥ stock_threshold_pct × COALESCE(NULLIF(ACS_D,0), default_acs_d)", italic=True)
    add_table(doc, ["Order", "Branch", "Condition", "Result"], OPT_TYPE_TABLE)

    add_heading(doc, "5.2 MIX aggregation (Part 3.7)", level=2)
    add_para(doc, "MIX rows are always aggregated. Legacy aliases: aggregate → st_maj, mark → each.")
    add_table(doc, ["mix_mode", "Behaviour"], MIX_MODES)

    add_heading(doc, "5.3 Sequencing & ordering", level=2)
    add_para(doc, "Working/alloc rows ordered: ST_RANK → MAJ_CAT → OPT_TYPE (RL=1, TBC=2, TBL=3) → OPT_PRIORITY_RANK → WERKS [→ SZ for alloc]. ST_RANK is per-MAJ_CAT priority rank of the store.")

    add_heading(doc, "5.4 OPT_MBQ rules", level=2)
    add_bullet(doc, "HOLD_DAYS extra days apply only to OPT_TYPE='TBL' (OPT_MBQ_WH).")
    add_bullet(doc, "OPT_MBQ = 0 when OPT_TYPE = 'MIX'.")
    add_bullet(doc, "For new articles (IS_NEW=1 AND AGE < age_threshold), use PER_OPT_SALE rate instead of ACS_D (age_threshold default 15).")

    add_heading(doc, "5.5 MJ_REQ caps (per OPT_TYPE downward cap)", level=2)
    add_para(doc, "After waterfall, SUM(SHIP_QTY) for each OPT_TYPE per (WERKS, MAJ_CAT) is clamped to cap_pct% × MJ_REQ. Defaults: rl_mj_req_cap_pct = tbc_mj_req_cap_pct = tbl_mj_req_cap_pct = 100.0. 0 disables.")

    add_heading(doc, "5.6 MBQ caps (Decision 4-B)", level=2)
    add_para(doc, "Anchored to MJ_MBQ_ORIG (NOT post-growth MJ_MBQ_REV). Defaults 100 = no over-ship vs original MBQ. 0 disables.")

    add_heading(doc, "5.7 Growth headroom (mj_req_growth_pct)", level=2)
    add_bullet(doc, "100 (default) = strict, waterfall stops at MAJ_CAT target.")
    add_bullet(doc, ">100 scales MJ_MBQ → MJ_MBQ_REV (sibling column; original preserved). MJ_REQ_REV = max(0, MJ_MBQ_REV − MJ_STK_TTL). MJ_REQ_ORIG retained for audit.")

    add_heading(doc, "5.8 Primary-grid coverage gate", level=2)
    add_bullet(doc, "pri_ct_check_rl = False (default): RL allowed even if PRI_CT% < 100 (boosted MBQ-cap path activates).")
    add_bullet(doc, "pri_ct_check_tbc = False (default): same for TBC.")
    add_bullet(doc, "TBL always enforces PRI_CT% ≥ 100 (R06 + revalidation SKIP_PRI_BROKEN).")

    add_heading(doc, "5.9 R07 size-coverage gate (TBL only)", level=2)
    add_para(doc, "Skip a TBL row when VAR_FNL_COUNT / VAR_COUNT < size_threshold AND VAR_FNL_COUNT < min_size_count. size_threshold defaults 0.6 — independent of stock_threshold_pct.")

    add_heading(doc, "5.10 Post-alloc OPT_STATUS (Part 8.5)", level=2)
    add_table(doc, ["OPT_TYPE", "ALLOC_QTY", "Post-alloc OPT_STATUS"], OPT_STATUS_TABLE)
    add_para(doc, "Also: TBL_LISTED_DATE = GETDATE() when OPT_TYPE='TBL' AND ALLOC_QTY > 0.", italic=True)

    # 6. Allocation
    add_heading(doc, "6. Allocation rules (per-OPT engine)", level=1)
    add_para(doc, "Source: backend/app/services/rule_engine_per_opt.py, rule_engine_new.py, rule_engine_pandas.py", italic=True)

    add_heading(doc, "6.1 Mode selection", level=2)
    add_table(doc, ["allocation_mode", "What runs"], ALLOC_MODES)
    add_para(doc, "exec_order = opt_type_first (default) = RL all rounds → TBC all rounds → TBL all rounds. round_first (R1 across all → R2 …) is configured but wiring still maps to opt_type_first today — see gap 9.2.")

    add_heading(doc, "6.2 Per-size dispatch formula (RL / TBC)", level=2)
    add_para(doc, "For each size of an OPT_TYPE in ('RL','TBC'), each round r:")
    formula = (
        "need_ship   = max(r × SZ_MBQ − SZ_STK − SHIP_QTY_so_far, 0)\n"
        "need_pool   = r × SZ_MBQ − SZ_STK − POOL_CONSUMED        (need_pool ← 0 when need_ship = 0)\n"
        "\n"
        "# Source priority: HOLD first, then MSA pool\n"
        "from_hold   = min(opt_need, RL_HOLD_QTY_remaining)        # per (WERKS, VAR_ART, SZ)\n"
        "opt_need   -= from_hold\n"
        "take_pool   = min(opt_need, live_pool)\n"
        "\n"
        "# Per-size ship ceiling (commits 2a8c87b + 58ef3dc, June 2026)\n"
        "ship_ceiling   = ceil(need_ship / pak) × pak\n"
        "raw_ship       = min(take_pool + from_hold, ship_ceiling)\n"
        "effective_ship = floor(raw_ship / pak) × pak   (or raw_ship when combined_supply = pak-multiple)\n"
        "pool_used      = max(effective_ship − from_hold, 0)"
    )
    code_para = doc.add_paragraph()
    code_run = code_para.add_run(formula)
    code_run.font.name = "Consolas"
    code_run.font.size = Pt(9)
    add_para(doc, "Cap guarantee: Σ SHIP per size ≤ I_ROD × SZ_MBQ − SZ_STK (plus up to pak−1 from pak rounding).", bold=True)

    add_heading(doc, "6.3 RL_HOLD_QTY semantics", level=2)
    add_para(doc, "RL_HOLD_QTY is NOT in-transit stock. It is store-specific stock reserved/earmarked at the RDC for a (WERKS, VAR_ART, SZ), sourced from ARS_NL_TBL_HOLD_TRACKING.HOLD_REM where IS_CLOSED=0 AND HOLD_REM>0. The engine consumes from the reservation first, then taps the MSA pool — total ship per size capped by the ship ceiling.")

    add_heading(doc, "6.4 TBL targeting", level=2)
    add_para(doc, "TBL ship target = SZ_MBQ_WH + (I_ROD − 1) × SZ_MBQ − SZ_STK   (hold counted once)")
    add_para(doc, "TBL SHIP and HOLD are pak-aligned independently (commit 58ef3dc).")

    add_heading(doc, "6.5 FNL_Q_REM semantics (per-OPT mode)", level=2)
    add_para(doc, "FNL_Q_REM on alloc rows is the live pool AFTER that OPT's draw. Pre-band snapshots in rule_engine_pandas and post-loop SQL recompute are gated off when per-OPT mode is on.")
    add_para(doc, "Audit pattern — combine FNL_Q_REM with ALLOC_REMARKS:")
    add_table(doc, ["FNL_Q_REM", "ALLOC_REMARKS excerpt", "Diagnosis"], FNL_Q_REM_TABLE)

    add_heading(doc, "6.6 Revalidation invariants", level=2)
    add_bullet(doc, "SKIP_PRI_BROKEN revalidation fires when an OPT lands without satisfying the active primary-grid gate.")
    add_bullet(doc, "Sec-cap dims must still be on every row — revalidation reads them.")
    add_bullet(doc, "ALLOC_REMARKS is the audit trail; never overwrite without preserving prior reasons.")

    # 7. Hold / Pend
    add_heading(doc, "7. Hold / Pending Allocation rules", level=1)

    add_heading(doc, "7.1 Lifecycles", level=2)
    add_bullet(doc, "PEND: queued → approved → dispatched | cancelled (reverted)")
    add_bullet(doc, "HOLD: active → parked (with reason) → unparked (back to active) | finalized")
    add_para(doc, "parked_history is the source of truth for HOLD transitions — every state change writes a row.")

    add_heading(doc, "7.2 Revert correctness", level=2)
    add_para(doc, "A reverted allocation must restore PEND to its pre-allocation value at the OPT grain. Recent commits flagging active correctness area: fecb6f9, 3c7693b, c06d051 — read git log before editing.")

    add_heading(doc, "7.3 Hold unpark must preserve OPT identity", level=2)
    add_para(doc, "A parked OPT must not re-enter active flow with a different OPT_TYPE (invariant 1) without an explicit transition row. Unparked rows must carry FAB / MACRO_MVGR / MICRO_MVGR / M_VND_CD / RNG_SEG back to alloc (invariant 4).")

    add_heading(doc, "7.4 Known correctness risks", level=2)
    add_table(doc, ["Risk", "Where", "Effect"], HOLD_PEND_RISKS)

    # 8. Checklist
    add_heading(doc, "8. Reviewer's checklist", level=1)
    add_para(doc, "Walk these in order — failure at step N makes step N+1 meaningless.")
    add_table(doc, ["Stage", "Check"], CHECKLIST)

    # 9. Gaps
    add_heading(doc, "9. Suggested rules / gaps spotted during this review", level=1)
    add_para(doc, "Points where the implementation, comments, and recorded rules disagree — bring to the next review meeting.")
    for gid, title, observation, suggestion in GAPS:
        add_heading(doc, f"{gid} {title}", level=2)
        add_para(doc, f"Observation: {observation}")
        add_para(doc, f"Suggested rule: {suggestion}", bold=True)

    # 10. Change log
    add_heading(doc, "10. Change log", level=1)
    add_table(doc, ["Date", "Change", "By"],
              [("2026-06-26", "Initial consolidated rule master created", "Santosh (drafted via Claude)")])
    add_para(doc, "How to extend: add new rules under the matching section. For module deep dives, also append a dated bullet to the relevant .claude/agents/ars_flow_kb/<area>.md.", italic=True)

    doc.save(out_path)


# ─── Excel (.xlsx) generation ─────────────────────────────────────────────────

HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")


def style_header_row(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = BORDER


def write_table(ws, start_row, headers, rows, widths=None):
    for c, h in enumerate(headers, start=1):
        ws.cell(row=start_row, column=c, value=h)
    style_header_row(ws, start_row, len(headers))
    for ri, row in enumerate(rows, start=start_row + 1):
        for ci, val in enumerate(row, start=1):
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.alignment = WRAP_TOP
            cell.border = BORDER
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[start_row].height = 22
    return start_row + 1 + len(rows)


def add_section_title(ws, row, text, ncols=4):
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(bold=True, size=12, color="1F4E78")
    cell.fill = SECTION_FILL
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    ws.row_dimensions[row].height = 22
    return row + 1


def build_xlsx(out_path: Path):
    wb = Workbook()

    # ── Sheet 0: Overview / TOC ──
    ws = wb.active
    ws.title = "Overview"
    ws["A1"] = "ARS Rule Master"
    ws["A1"].font = Font(bold=True, size=18, color="1F4E78")
    ws["A2"] = "MSA · Grid Builder · Listing · Allocation · Hold/Pending Allocation"
    ws["A2"].font = Font(italic=True, size=11)
    ws["A3"] = "Generated 2026-06-26 from docs/RULE_MASTER.md"
    ws["A3"].font = Font(italic=True, size=10, color="606060")

    toc = [
        ("Invariants",       "Six cross-cutting invariants spanning all stages"),
        ("MSA",              "11-step flow, universe rule, 5 reconciliation guarantees, 8 write paths"),
        ("Grid",             "Taxonomy, growth %, sec-cap dispatch, schema management"),
        ("Listing",          "OPT_TYPE decision table, MIX modes, caps, growth headroom, gates"),
        ("Allocation",       "Per-OPT engine, ship ceiling, hold-first, FNL_Q_REM semantics"),
        ("Hold-Pend",        "Lifecycles, revert correctness, 7 known correctness risks"),
        ("Checklist",        "Stage-by-stage reviewer audit ladder"),
        ("Gaps",             "Suggested new rules / inconsistencies found while writing"),
    ]
    ws["A5"] = "Sheet"
    ws["B5"] = "Contents"
    style_header_row(ws, 5, 2)
    for i, (sheet, desc) in enumerate(toc, start=6):
        ws.cell(row=i, column=1, value=sheet).alignment = WRAP_TOP
        ws.cell(row=i, column=2, value=desc).alignment = WRAP_TOP
        ws.cell(row=i, column=1).border = BORDER
        ws.cell(row=i, column=2).border = BORDER
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 90

    # ── Invariants ──
    ws = wb.create_sheet("Invariants")
    write_table(ws, 1, ["#", "Invariant", "Meaning", "Common violation site"], INVARIANTS,
                widths=[6, 32, 60, 50])

    # ── MSA ──
    ws = wb.create_sheet("MSA")
    r = add_section_title(ws, 1, "3.1  11-step MSA flow (universe-anchored)", ncols=3)
    r = write_table(ws, r, ["Step", "Action", "Rule"], MSA_STEPS, widths=[8, 28, 90])
    r += 1
    r = add_section_title(ws, r, "3.3  Reconciliation guarantees (must all hold)", ncols=3)
    r = write_table(ws, r, ["ID", "Check", "Pass condition"], MSA_RECON, widths=[8, 90, 18])
    r += 1
    r = add_section_title(ws, r, "3.4  Eight write paths to ARS_MSA_*", ncols=4)
    r = write_table(ws, r, ["#", "Function", "Trigger", "Writes"], MSA_WRITE_PATHS,
                    widths=[6, 38, 48, 48])
    r += 1
    r = add_section_title(ws, r, "3.2  Universe rule", ncols=3)
    note = ("Stock contribution stays SLOC-scoped. Products with stock only on non-selected shelves do NOT enter MSA. "
            "The universe expands ONLY when a real obligation (PEND or HOLD) attaches to the (RDC, GEN_ART) pair. "
            "Cross-shelf stock alone is not enough.")
    cell = ws.cell(row=r, column=1, value=note)
    cell.alignment = WRAP_TOP
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    ws.row_dimensions[r].height = 55

    # ── Grid ──
    ws = wb.create_sheet("Grid")
    r = add_section_title(ws, 1, "4.1  Grid taxonomy", ncols=4)
    r = write_table(ws, r, ["Grid", "Type", "Grain", "Purpose"], GRID_TAXONOMY,
                    widths=[20, 14, 28, 50])
    r += 1
    r = add_section_title(ws, r, "4.5  Grid schema management rules", ncols=2)
    r = write_table(ws, r, ["Rule", "Reason"], GRID_SCHEMA_RULES, widths=[50, 70])
    r += 1
    r = add_section_title(ws, r, "4.4  Sec-cap dispatch (strict mode)", ncols=2)
    sec_lines = [
        ("Active when",         "apply_sec_cap_in_normal = True (default since 2026-06-16)"),
        ("Per-grid trigger",    "ARS_GRID_BUILDER.sec_cap_applicable = 1"),
        ("Budget formula",      "budget = max(0, MBQ_ORIG × sec_cap_pct% − STK_TTL)"),
        ("Breach behaviour",    "Block the OPT row at the main-pass pre-gate"),
        ("Per-grid %",          "ARS_GRID_BUILDER.sec_cap_pct (default 130 when NULL)"),
        ("High-demand override","REMOVED — strict binary semantics only"),
    ]
    r = write_table(ws, r, ["Field", "Value / behaviour"], sec_lines, widths=[24, 70])

    # ── Listing ──
    ws = wb.create_sheet("Listing")
    r = add_section_title(ws, 1, "5.1  OPT_TYPE classification (Part 3.6) — decision table", ncols=4)
    r = write_table(ws, r, ["Order", "Branch", "Condition", "Result"], OPT_TYPE_TABLE,
                    widths=[8, 38, 70, 10])
    r += 1
    r = add_section_title(ws, r, "Symbols used in conditions", ncols=2)
    r = write_table(ws, r, ["Symbol", "Definition"],
                    [("SUPPLY_OK",     "(MSA_FNL_Q > 0) OR (RL_HOLD_QTY > 0)"),
                     ("ADEQUATE_STK",  "STK_TTL ≥ stock_threshold_pct × COALESCE(NULLIF(ACS_D,0), default_acs_d)"),
                     ("Defaults",       "stock_threshold_pct=0.6, default_acs_d=18, min_size_count=3")],
                    widths=[20, 90])
    r += 1
    r = add_section_title(ws, r, "5.2  MIX aggregation modes", ncols=2)
    r = write_table(ws, r, ["mix_mode", "Behaviour"], MIX_MODES, widths=[24, 70])
    r += 1
    r = add_section_title(ws, r, "5.10  Post-alloc OPT_STATUS", ncols=3)
    r = write_table(ws, r, ["OPT_TYPE", "ALLOC_QTY", "Post-alloc OPT_STATUS"], OPT_STATUS_TABLE,
                    widths=[14, 22, 36])
    r += 1
    r = add_section_title(ws, r, "5.5 / 5.6  Caps (per OPT_TYPE)", ncols=4)
    cap_rows = [
        ("MJ_REQ cap",  "rl_mj_req_cap_pct",  "100.0", "SUM(SHIP_QTY) per OPT_TYPE per (WERKS, MAJ_CAT) ≤ cap_pct × MJ_REQ. 0 disables."),
        ("MJ_REQ cap",  "tbc_mj_req_cap_pct", "100.0", "Same as above for TBC."),
        ("MJ_REQ cap",  "tbl_mj_req_cap_pct", "100.0", "Same as above for TBL."),
        ("MBQ cap",     "rl_mbq_cap_pct",     "100.0", "Anchored to MJ_MBQ_ORIG (not post-growth REV)."),
        ("MBQ cap",     "tbc_mbq_cap_pct",    "100.0", "Same as above for TBC."),
        ("MBQ cap",     "tbl_mbq_cap_pct",    "100.0", "Same as above for TBL."),
        ("Growth",      "mj_req_growth_pct",  "100.0", "100 strict; >100 scales MJ_MBQ → MJ_MBQ_REV; MJ_REQ_REV = max(0, MJ_MBQ_REV − MJ_STK_TTL)."),
    ]
    r = write_table(ws, r, ["Cap type", "Parameter", "Default", "Effect"], cap_rows,
                    widths=[14, 24, 12, 80])
    r += 1
    r = add_section_title(ws, r, "5.8  Primary-grid coverage gate", ncols=2)
    r = write_table(ws, r, ["Toggle", "Behaviour"],
                    [("pri_ct_check_rl = False (default)",   "RL allowed even if PRI_CT% < 100 (boosted MBQ-cap path)"),
                     ("pri_ct_check_tbc = False (default)",  "Same for TBC"),
                     ("TBL",                                  "Always enforces PRI_CT% ≥ 100 (R06 + SKIP_PRI_BROKEN)")],
                    widths=[36, 70])

    # ── Allocation ──
    ws = wb.create_sheet("Allocation")
    r = add_section_title(ws, 1, "6.1  Allocation modes", ncols=2)
    r = write_table(ws, r, ["allocation_mode", "What runs"], ALLOC_MODES, widths=[24, 70])
    r += 1
    r = add_section_title(ws, r, "6.2  Per-size dispatch formula (RL / TBC)", ncols=2)
    formula_rows = [
        ("need_ship",     "max(r × SZ_MBQ − SZ_STK − SHIP_QTY_so_far, 0)"),
        ("need_pool",     "r × SZ_MBQ − SZ_STK − POOL_CONSUMED  (need_pool ← 0 when need_ship = 0)"),
        ("from_hold",     "min(opt_need, RL_HOLD_QTY_remaining)  per (WERKS, VAR_ART, SZ)"),
        ("opt_need",      "opt_need − from_hold"),
        ("take_pool",     "min(opt_need, live_pool)"),
        ("ship_ceiling",  "ceil(need_ship / pak) × pak"),
        ("raw_ship",      "min(take_pool + from_hold, ship_ceiling)"),
        ("effective_ship","floor(raw_ship / pak) × pak  (or raw_ship when combined_supply is pak-multiple)"),
        ("pool_used",     "max(effective_ship − from_hold, 0)"),
        ("Cap guarantee", "Σ SHIP per size ≤ I_ROD × SZ_MBQ − SZ_STK (+ up to pak−1 from pak rounding)"),
    ]
    r = write_table(ws, r, ["Term", "Formula"], formula_rows, widths=[20, 90])
    r += 1
    r = add_section_title(ws, r, "6.5  FNL_Q_REM × ALLOC_REMARKS audit pattern", ncols=3)
    r = write_table(ws, r, ["FNL_Q_REM", "ALLOC_REMARKS excerpt", "Diagnosis"], FNL_Q_REM_TABLE,
                    widths=[14, 42, 50])
    r += 1
    r = add_section_title(ws, r, "6.3 / 6.4  Hold semantics & TBL target", ncols=2)
    r = write_table(ws, r, ["Topic", "Rule"],
                    [("RL_HOLD_QTY",         "NOT in-transit. Store-specific stock reserved/earmarked at RDC for (WERKS, VAR_ART, SZ) from ARS_NL_TBL_HOLD_TRACKING.HOLD_REM. Consumed BEFORE MSA pool."),
                     ("TBL ship target",     "SZ_MBQ_WH + (I_ROD − 1) × SZ_MBQ − SZ_STK  (hold counted once)"),
                     ("TBL pak alignment",   "TBL SHIP and HOLD are pak-aligned independently (commit 58ef3dc).")],
                    widths=[22, 90])

    # ── Hold / Pend ──
    ws = wb.create_sheet("Hold-Pend")
    r = add_section_title(ws, 1, "7.1  Lifecycles", ncols=2)
    r = write_table(ws, r, ["Type", "Lifecycle"],
                    [("PEND", "queued → approved → dispatched | cancelled (reverted)"),
                     ("HOLD", "active → parked (with reason) → unparked (back to active) | finalized")],
                    widths=[12, 90])
    r += 1
    r = add_section_title(ws, r, "7.4  Known correctness risks", ncols=3)
    r = write_table(ws, r, ["Risk", "Where", "Effect"], HOLD_PEND_RISKS,
                    widths=[48, 32, 60])

    # ── Checklist ──
    ws = wb.create_sheet("Checklist")
    r = add_section_title(ws, 1, "8.  Reviewer's checklist (walk top-to-bottom)", ncols=3)
    headers = ["Stage", "Check", "Pass / Fail / N-A"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=r, column=c, value=h)
    style_header_row(ws, r, len(headers))
    r += 1
    for stage, item in CHECKLIST:
        ws.cell(row=r, column=1, value=stage).alignment = WRAP_TOP
        ws.cell(row=r, column=2, value=item).alignment = WRAP_TOP
        ws.cell(row=r, column=3, value="").alignment = WRAP_TOP
        for c in range(1, 4):
            ws.cell(row=r, column=c).border = BORDER
        r += 1
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 90
    ws.column_dimensions["C"].width = 24

    # ── Gaps ──
    ws = wb.create_sheet("Gaps")
    r = add_section_title(ws, 1, "9.  Suggested rules / gaps spotted during this review", ncols=4)
    r = write_table(ws, r, ["ID", "Title", "Observation", "Suggested rule"],
                    GAPS, widths=[8, 40, 70, 60])

    wb.save(out_path)


# ─── Entry ────────────────────────────────────────────────────────────────────

def main():
    here = Path(__file__).resolve().parent
    docx_path = here / "RULE_MASTER.docx"
    xlsx_path = here / "RULE_MASTER.xlsx"
    build_docx(docx_path)
    build_xlsx(xlsx_path)
    print(f"WROTE  {docx_path}  ({docx_path.stat().st_size:,} bytes)")
    print(f"WROTE  {xlsx_path}  ({xlsx_path.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
